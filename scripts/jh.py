#!/usr/bin/env python3
# /// script
# requires-python = ">=3.9"
# dependencies = [
#     "fpdf2>=2.8,<3",
#     "python-docx>=1.1,<2",
#     "pypdf>=5,<7",
#     "jsonschema>=4.18,<5",
# ]
# ///
"""jh: helper CLI for the job-hunting Claude Code plugin.

Run it with `uv run jh.py <command>`; uv installs the dependencies listed above on
first use. Commands print JSON to stdout (except `extract` and `history list`, which
print text) so the calling skill can read the result.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import html
import io
import json
import logging
import math
import os
import platform
import re
import shutil
import sys
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
from html.parser import HTMLParser
from pathlib import Path

PLUGIN_ROOT = Path(__file__).resolve().parent.parent
SCHEMA_DIR = PLUGIN_ROOT / "schemas"
DEFAULT_DATA_DIR = "~/job-hunting"
STATUSES = ("tailored", "applied", "interviewing", "offer", "rejected", "withdrawn")
PT = 0.3528  # millimetres per point

logging.getLogger("fontTools").setLevel(logging.ERROR)
logging.getLogger("fpdf").setLevel(logging.ERROR)  # missing glyphs are reported in the JSON result instead


# ---------------------------------------------------------------------------
# Small helpers


def emit(obj) -> None:
    print(json.dumps(obj, ensure_ascii=False, indent=2))


def die(message: str, **extra) -> None:
    emit({"ok": False, "error": message, **extra})
    sys.exit(1)


def now_iso() -> str:
    return dt.datetime.now().astimezone().isoformat(timespec="seconds")


def read_json(path):
    try:
        return json.loads(Path(path).read_text(encoding="utf-8-sig"))
    except FileNotFoundError:
        die(f"File not found: {path}")
    except json.JSONDecodeError as e:
        die(f"Invalid JSON in {path}: {e}")


def write_json(path, data) -> None:
    Path(path).write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def resolve_data_dir(raw) -> Path:
    """Resolve the data folder; an empty or unsubstituted setting falls back to the default."""
    value = (raw or "").strip().strip("'\"")
    if not value or "${" in value:
        value = os.environ.get("JOB_HUNTING_DATA_DIR") or DEFAULT_DATA_DIR
    return Path(os.path.expandvars(os.path.expanduser(value))).resolve()


def safe_name(text: str, max_len: int = 60) -> str:
    """File-system safe name that keeps non-ASCII letters such as CJK company names."""
    text = unicodedata.normalize("NFKC", text or "").strip()
    text = re.sub(r"[^\w.-]+", "-", text).strip("-.")
    return text[:max_len].rstrip("-.") or "untitled"


def iter_strings(obj):
    if isinstance(obj, str):
        yield obj
    elif isinstance(obj, dict):
        for value in obj.values():
            yield from iter_strings(value)
    elif isinstance(obj, list):
        for value in obj:
            yield from iter_strings(value)


def item_text(item) -> str:
    return item.get("text", "") if isinstance(item, dict) else str(item)


def split_paragraphs(text: str):
    return [p.strip() for p in (text or "").split("\n") if p.strip()]


def section_has_content(sec) -> bool:
    kind = sec.get("kind")
    if kind == "text":
        return bool(split_paragraphs(sec.get("text", "")))
    if kind == "entries":
        return any(e.get("heading") or e.get("bullets") or e.get("summary") for e in sec.get("entries", []))
    if kind == "groups":
        return any(g.get("items") for g in sec.get("groups", []))
    if kind == "list":
        return any(item_text(i).strip() for i in sec.get("items", []))
    return False


def normalize(text: str) -> str:
    text = unicodedata.normalize("NFKC", text or "")
    text = re.sub(r"[\u2010-\u2015\u2212]", "-", text)
    return re.sub(r"\s+", " ", text).strip().lower()


HAN_RE = re.compile(r"[\u2e80-\u2fdf\u3000-\u303f\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff\uff00-\uffef]")
KANA_RE = re.compile(r"[\u3040-\u30ff\u31f0-\u31ff]")
HANGUL_RE = re.compile(r"[\u1100-\u11ff\u3130-\u318f\uac00-\ud7af]")


def keyword_in(normalized_text: str, term: str) -> bool:
    needle = normalize(term)
    if not needle:
        return False
    if HAN_RE.search(needle) or KANA_RE.search(needle) or HANGUL_RE.search(needle):
        return needle in normalized_text
    return re.search(r"(?<![0-9a-z])" + re.escape(needle) + r"(?![0-9a-z])", normalized_text) is not None


def schema_errors(data, kind: str):
    import jsonschema

    schema = json.loads((SCHEMA_DIR / f"{kind}.schema.json").read_text(encoding="utf-8"))
    validator = jsonschema.Draft202012Validator(schema)
    errors = sorted(validator.iter_errors(data), key=lambda e: [str(p) for p in e.absolute_path])
    return [f"{'/'.join(str(p) for p in e.absolute_path) or '(root)'}: {e.message}" for e in errors]


def collect_ids(obj, acc=None):
    acc = [] if acc is None else acc
    if isinstance(obj, dict):
        if isinstance(obj.get("id"), str):
            acc.append(obj["id"])
        for value in obj.values():
            collect_ids(value, acc)
    elif isinstance(obj, list):
        for value in obj:
            collect_ids(value, acc)
    return acc


# ---------------------------------------------------------------------------
# Fonts


# The second entry finds the Windows fonts from WSL.
WIN_FONT_DIRS = (Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts", Path("/mnt/c/Windows/Fonts"))
MAC_FONTS = Path("/System/Library/Fonts")
FC_LANG = {"zh-Hant": "zh-tw", "zh-Hans": "zh-cn", "ja": "ja", "ko": "ko"}
NOTO_CJK_DIRS = (
    "/usr/share/fonts/opentype/noto",
    "/usr/share/fonts/google-noto-cjk",
    "/usr/share/fonts/noto-cjk",
    "/usr/share/fonts/OTF",
)
NOTO_CJK_INDEX = {"ja": 0, "ko": 1, "zh-Hans": 2, "zh-Hant": 3}
DOCX_EAST_ASIAN_FONT = {
    "zh-Hant": "Microsoft JhengHei",
    "zh-Hans": "Microsoft YaHei",
    "ja": "Yu Gothic",
    "ko": "Malgun Gothic",
}


def fc_match(pattern: str):
    """Ask fontconfig (Linux, some macOS setups) for a font file; returns (path, index) or None."""
    if not shutil.which("fc-match"):
        return None
    import subprocess

    try:
        out = subprocess.run(["fc-match", "-f", "%{file}|%{index}", pattern], capture_output=True,
                             text=True, timeout=10).stdout.strip()
    except (OSError, subprocess.SubprocessError):
        return None
    path, _, index = out.partition("|")
    if Path(path).suffix.lower() in (".ttf", ".ttc", ".otf") and Path(path).exists():
        return path, int(index or 0)
    return None


def latin_font_candidates():
    sets = [(w / "arial.ttf", w / "arialbd.ttf", w / "ariali.ttf", w / "arialbi.ttf") for w in WIN_FONT_DIRS]
    for d in map(Path, ("/System/Library/Fonts/Supplemental", "/Library/Fonts")):
        sets.append((d / "Arial.ttf", d / "Arial Bold.ttf", d / "Arial Italic.ttf", d / "Arial Bold Italic.ttf"))
    for d in map(Path, (
        "/usr/share/fonts/truetype/liberation",
        "/usr/share/fonts/truetype/liberation2",
        "/usr/share/fonts/liberation-sans",
        "/usr/share/fonts/liberation-sans-fonts",
        "/usr/share/fonts/liberation",
        "/usr/share/fonts/TTF",
    )):
        sets.append((d / "LiberationSans-Regular.ttf", d / "LiberationSans-Bold.ttf",
                     d / "LiberationSans-Italic.ttf", d / "LiberationSans-BoldItalic.ttf"))
    for d in map(Path, ("/usr/share/fonts/truetype/dejavu", "/usr/share/fonts/dejavu-sans-fonts",
                        "/usr/share/fonts/dejavu", "/usr/share/fonts/TTF")):
        sets.append((d / "DejaVuSans.ttf", d / "DejaVuSans-Bold.ttf",
                     d / "DejaVuSans-Oblique.ttf", d / "DejaVuSans-BoldOblique.ttf"))
    return sets


def find_latin_font(override=None):
    if override:
        path = Path(override).expanduser()
        if not path.exists():
            die(f"Font not found: {path}")
        return {"regular": str(path), "bold": None, "italic": None, "bolditalic": None}
    for regular, bold, italic, bold_italic in latin_font_candidates():
        if regular.exists():
            return {
                "regular": str(regular),
                "bold": str(bold) if bold.exists() else None,
                "italic": str(italic) if italic.exists() else None,
                "bolditalic": str(bold_italic) if bold_italic.exists() else None,
            }
    # fontconfig fallback; TrueType collections are skipped because add_font needs an index per style.
    found = {key: fc_match(pattern) for key, pattern in (
        ("regular", "sans-serif"), ("bold", "sans-serif:bold"), ("italic", "sans-serif:italic"),
        ("bolditalic", "sans-serif:bold:italic"))}
    if found["regular"] and not found["regular"][0].lower().endswith(".ttc"):
        return {key: (value[0] if value and not value[0].lower().endswith(".ttc") else None)
                for key, value in found.items()}
    return None


def cjk_font_candidates(script: str):
    m = MAC_FONTS
    table = []
    for w in WIN_FONT_DIRS:
        table += {
            "zh-Hant": [(w / "msjh.ttc", 0, w / "msjhbd.ttc", 0), (w / "mingliu.ttc", 0, w / "mingliub.ttc", 0)],
            "zh-Hans": [(w / "msyh.ttc", 0, w / "msyhbd.ttc", 0), (w / "simsun.ttc", 0, None, 0)],
            "ja": [(w / "YuGothM.ttc", 0, w / "YuGothB.ttc", 0), (w / "meiryo.ttc", 0, w / "meiryob.ttc", 0),
                   (w / "msgothic.ttc", 0, None, 0)],
            "ko": [(w / "malgun.ttf", 0, w / "malgunbd.ttf", 0)],
        }[script]
    table += {
        "zh-Hant": [(m / "PingFang.ttc", 0, None, 0)],
        "zh-Hans": [(m / "PingFang.ttc", 0, None, 0)],
        "ja": [(m / "ヒラギノ角ゴシック W3.ttc", 0, m / "ヒラギノ角ゴシック W6.ttc", 0)],
        "ko": [(m / "AppleSDGothicNeo.ttc", 0, None, 0)],
    }[script]
    index = NOTO_CJK_INDEX[script]
    for d in map(Path, NOTO_CJK_DIRS):
        table.append((d / "NotoSansCJK-Regular.ttc", index, d / "NotoSansCJK-Bold.ttc", index))
    return table


def find_cjk_font(script: str):
    for regular, regular_index, bold, bold_index in cjk_font_candidates(script):
        if regular.exists():
            return (str(regular), regular_index, str(bold) if bold and bold.exists() else None, bold_index)
    found = fc_match(f"sans-serif:lang={FC_LANG[script]}")
    return (found[0], found[1], None, 0) if found else None


def preferred_han_script(language: str, market: str) -> str:
    lang = (language or "").lower()
    if lang.startswith("ja"):
        return "ja"
    if lang.startswith("ko"):
        return "ko"
    if lang.startswith("zh"):
        return "zh-Hant" if any(tag in lang for tag in ("hant", "-tw", "-hk", "-mo")) else "zh-Hans"
    return {"HK": "zh-Hant", "TW": "zh-Hant", "MO": "zh-Hant", "CN": "zh-Hans", "SG": "zh-Hans",
            "MY": "zh-Hans", "JP": "ja", "KR": "ko"}.get((market or "").upper(), "zh-Hant")


def scripts_needed(text: str, language: str, market: str):
    """CJK scripts the text needs, preferred first. Han falls back to other Han-capable fonts."""
    needed = []
    if HAN_RE.search(text):
        preferred = preferred_han_script(language, market)
        needed.append([preferred] + [s for s in ("zh-Hant", "zh-Hans", "ja") if s != preferred])
    if KANA_RE.search(text):
        needed.append(["ja"])
    if HANGUL_RE.search(text):
        needed.append(["ko"])
    return needed


CORE_REPLACEMENTS = {"\u2022": "-", "\u2013": "-", "\u2014": "-", "\u2018": "'", "\u2019": "'",
                     "\u201c": '"', "\u201d": '"', "\u2026": "...", "\u00a0": " "}


def core_safe(text: str) -> str:
    text = "".join(CORE_REPLACEMENTS.get(ch, ch) for ch in text or "")
    return text.encode("latin-1", "replace").decode("latin-1")


# Scripts that need text shaping or right-to-left layout, which the PDF renderer does not do.
UNSHAPED_SCRIPTS = {
    "Hebrew": "֐-׿יִ-ﭏ",
    "Arabic": "؀-ۿݐ-ݿࢠ-ࣿﭐ-﷿ﹰ-﻿",
    "Syriac, Thaana or N'Ko": "܀-߿",
    "Devanagari or another Indic script": "ऀ-෿",
    "Thai": "฀-๿",
    "Lao": "຀-໿",
    "Tibetan": "ༀ-࿿",
    "Myanmar": "က-႟",
    "Khmer": "ក-៿",
}


def unshaped_scripts(text: str):
    return [name for name, chars in UNSHAPED_SCRIPTS.items() if re.search(f"[{chars}]", text)]


def register_fonts(pdf, text, language, market, font=None, cjk_font=None):
    """Register the body font plus CJK fallbacks. Returns (family, styles, core_only, warnings)."""
    warnings = []
    latin = find_latin_font(font)
    if not latin:
        unsupported = sorted({ch for ch in text if ch not in CORE_REPLACEMENTS and ord(ch) > 0xFF})
        if unsupported:
            die("No TrueType font was found, so these characters cannot be rendered: " + "".join(unsupported[:40])
                + ". Install a font such as Liberation Sans or DejaVu Sans (and Noto Sans CJK for CJK text), "
                "or pass --font and --cjk-font.")
        warnings.append("No Arial/Liberation/DejaVu font found; used Helvetica (Latin-1 only).")
        return "helvetica", {"", "B", "I", "BI"}, True, warnings

    pdf.add_font("Body", "", latin["regular"])
    styles = {""}
    for style, key in (("B", "bold"), ("I", "italic"), ("BI", "bolditalic")):
        if latin[key]:
            pdf.add_font("Body", style, latin[key])
            styles.add(style)

    fallbacks = []
    for options in scripts_needed(text, language, market):
        if cjk_font:
            path, _, index = cjk_font.partition("#")
            found = (path, int(index or 0), None, 0)
        else:
            found = next((f for f in map(find_cjk_font, options) if f), None)
        if not found:
            warnings.append(f"No font found for {options[0]} text; those characters may be missing. Pass --cjk-font.")
            continue
        name = "CJK" + str(len(fallbacks))
        try:
            pdf.add_font(name, "", found[0], collection_font_number=found[1])
            if found[2]:
                pdf.add_font(name, "B", found[2], collection_font_number=found[3])
            fallbacks.append(name)
        except Exception as e:  # unsupported font file: keep going with what we have
            warnings.append(f"Could not load {found[0]}: {e}")
        if cjk_font:
            break
    if fallbacks:
        pdf.set_fallback_fonts(fallbacks, exact_match=False)
    return "Body", styles, False, warnings


# ---------------------------------------------------------------------------
# Text extraction


def pdf_pages_text(path: Path):
    from pypdf import PdfReader

    return [page.extract_text() or "" for page in PdfReader(str(path)).pages]


def docx_text(path: Path) -> str:
    """Body, header/footer and text-box text of a DOCX in reading order."""
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(str(path))
    w_p, w_t, w_tab, w_br = qn("w:p"), qn("w:t"), qn("w:tab"), qn("w:br")
    fallback = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"

    def para_text(p):
        parts = []
        for node in p.iter(w_t, w_tab, w_br):
            ancestor = node.getparent()
            while ancestor is not None and ancestor is not p:
                if ancestor.tag in (w_p, fallback):
                    break
                ancestor = ancestor.getparent()
            if ancestor is not p:
                continue  # belongs to a nested text-box paragraph or a compatibility copy
            if node.tag == w_t:
                parts.append(node.text or "")
            else:
                parts.append("\t" if node.tag == w_tab else "\n")
        return "".join(parts).strip()

    def lines_of(root):
        lines = []
        for p in root.iter(w_p):
            if any(a.tag == fallback for a in p.iterancestors()):
                continue
            text = para_text(p)
            if text:
                lines.append(text)
        return lines

    out = []
    for section in doc.sections:
        for part, label in ((section.header, "header"), (section.footer, "footer")):
            if part.is_linked_to_previous:
                continue
            lines = lines_of(part._element)
            if lines:
                out.append(f"[{label}] " + " | ".join(lines))
    out.extend(lines_of(doc.element.body))
    return "\n".join(out)


def extract_text(path: Path):
    """Return (text, page_count or None)."""
    suffix = path.suffix.lower()
    if not path.exists():
        die(f"File not found: {path}")
    if suffix == ".pdf":
        pages = pdf_pages_text(path)
        return "\n\n".join(f"--- page {i} ---\n{t}" for i, t in enumerate(pages, 1)), len(pages)
    if suffix == ".docx":
        return docx_text(path), None
    if suffix in (".txt", ".md"):
        return path.read_text(encoding="utf-8-sig", errors="replace"), None
    die(f"Unsupported file type {suffix!r}. Ask the user to export the CV as DOCX or PDF.")


# ---------------------------------------------------------------------------
# PDF rendering


EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
DOMAIN_RE = re.compile(r"^(www\.)?[a-z0-9-]+(\.[a-z0-9-]+)*\.[a-z]{2,}(/\S*)?$", re.I)


def contact_link(item: str) -> str:
    s = item.strip()
    if EMAIL_RE.match(s):
        return "mailto:" + s
    if re.match(r"^https?://", s, re.I):
        return s
    if DOMAIN_RE.match(s):
        return "https://" + s
    return ""


# CJK characters are break opportunities; Latin words stay whole.
CJK_CHARS = "ᄀ-ᇿ⺀-〿぀-ヿ㄰-㆏ㇰ-ㇿ㐀-䶿一-鿿가-힯豈-﫿＀-￯"
CJK_ANY_RE = re.compile(f"[{CJK_CHARS}]")
LINE_TOKEN_RE = re.compile(f"[{CJK_CHARS}]|[^\\s{CJK_CHARS}]+\\s*|\\s+")
NO_LINE_START = set(")]}%,.;:!?’”…、。」』】〉》・ー），．：；！？")
NO_LINE_END = set("([{‘“「『【〈《（")


def line_tokens(text: str):
    tokens = []
    for token in LINE_TOKEN_RE.findall(text):
        if tokens and (token[0] in NO_LINE_START or tokens[-1].rstrip()[-1:] in NO_LINE_END):
            tokens[-1] += token
        else:
            tokens.append(token)
    return tokens


class PdfCV:
    PRIMARY = (31, 35, 40)
    SECONDARY = (87, 96, 106)
    RULE = (160, 160, 160)

    def __init__(self, cv, font=None, cjk_font=None):
        from fpdf import FPDF
        from fpdf.enums import XPos, YPos

        self.XPos, self.YPos = XPos, YPos
        self.cv = cv
        meta = cv.get("meta", {})
        self.lang = meta.get("language") or "en"
        self.base = float(meta.get("fontSize") or 10.5)
        paper = "letter" if str(meta.get("paper", "A4")).lower() == "letter" else "a4"
        self.pdf = FPDF(unit="mm", format=paper)
        side = 17.8 if paper == "letter" else 17.0
        self.pdf.set_margins(side, 15, side)
        self.pdf.set_auto_page_break(True, margin=15)
        text = "\n".join(iter_strings(cv))
        scripts = unshaped_scripts(text)
        if scripts:
            die(f"The PDF renderer cannot lay out {', '.join(scripts)} text correctly. Render a DOCX instead "
                "(--formats docx) or write those parts in Latin script.")
        self.missing_glyphs = []
        self.family, self.styles, self.core, self.warnings = register_fonts(
            self.pdf, text, self.lang, meta.get("market") or "", font, cjk_font)

    def t(self, text) -> str:
        return core_safe(text) if self.core else (text or "")

    def font(self, style="", size=None, color=None):
        if style not in self.styles:
            style = "B" if "B" in style and "B" in self.styles else ""
        self.pdf.set_font(self.family, style, size or self.base)
        self.pdf.set_text_color(*(color or self.PRIMARY))

    def lh(self, size=None) -> float:
        return (size or self.base) * PT * 1.32

    def ensure(self, needed: float):
        pdf = self.pdf
        if pdf.get_y() + needed > pdf.h - pdf.b_margin:
            pdf.add_page()

    def wrap(self, text: str, width: float):
        """Greedy line breaking in the current font."""
        pdf, lines, current = self.pdf, [], ""
        tokens = []
        for token in line_tokens(text):
            while pdf.get_string_width(token.rstrip()) > width and len(token) > 1:  # e.g. a long URL
                cut = len(token) - 1
                while cut > 1 and pdf.get_string_width(token[:cut]) > width:
                    cut -= 1
                tokens.append(token[:cut])
                token = token[cut:]
            tokens.append(token)
        for token in tokens:
            if not current.strip() or pdf.get_string_width((current + token).rstrip()) <= width:
                current += token
            else:
                lines.append(current.rstrip())
                current = token.lstrip()
        if current.strip():
            lines.append(current.rstrip())
        return lines

    def text_at(self, text, x, width, h):
        """Write wrapped text starting at x, then move to the start of the next line."""
        pdf, text = self.pdf, self.t(text)
        if not CJK_ANY_RE.search(text):
            pdf.set_x(x)
            pdf.multi_cell(width, h, text, align="L", new_x=self.XPos.LMARGIN, new_y=self.YPos.NEXT)
            return
        for line in self.wrap(text, width):
            pdf.set_x(x)
            pdf.cell(width, h, line, new_x=self.XPos.LMARGIN, new_y=self.YPos.NEXT)

    def block(self, text, size=None, style="", color=None, after=0.0):
        self.font(style, size, color)
        self.text_at(text, self.pdf.l_margin, self.pdf.epw, self.lh(size))
        if after:
            self.pdf.ln(after)

    def inline(self, items, links=False, color=None):
        size = self.base - 1
        self.font("", size, color)
        h = self.lh(size)
        for i, item in enumerate(items):
            if i:
                self.pdf.write(h, self.t("  |  "))
            self.pdf.write(h, self.t(item), link=contact_link(item) if links else "")
        self.pdf.ln(h)

    def header(self):
        h = self.cv["header"]
        self.block(h["name"], size=self.base + 7.5, style="B", after=0.6)
        if h.get("headline"):
            self.block(h["headline"], size=self.base + 0.5, color=self.SECONDARY, after=0.4)
        self.inline(h.get("contact", []), links=True)
        if h.get("details"):
            self.inline(h["details"], color=self.SECONDARY)
        self.pdf.ln(0.8)

    def section_title(self, title, first_block_lines):
        # Keep the title on the same page as the start of the section's first block.
        pdf, size = self.pdf, self.base + 1
        self.ensure(self.lh(size) + 4.4 + self.lh() * first_block_lines)
        pdf.ln(2.4)
        self.font("B", size)
        pdf.cell(0, self.lh(size), self.t(title), new_x=self.XPos.LMARGIN, new_y=self.YPos.NEXT)
        y = pdf.get_y() + 0.4
        pdf.set_draw_color(*self.RULE)
        pdf.set_line_width(0.25)
        pdf.line(pdf.l_margin, y, pdf.w - pdf.r_margin, y)
        pdf.set_y(y + 1.6)

    def bullet(self, text):
        pdf = self.pdf
        indent, gap = 1.6, 3.4
        self.ensure(self.lh() * 1.2)
        self.font("")
        pdf.set_x(pdf.l_margin + indent)
        pdf.cell(gap, self.lh(), self.t("\u2022"))
        self.text_at(text, pdf.l_margin + indent + gap, pdf.epw - indent - gap, self.lh())
        pdf.ln(0.5)

    def entry(self, e):
        pdf = self.pdf
        self.ensure(self.lh() * 4)
        dates = self.t(e.get("dates") or "")
        self.font("")
        dates_w = pdf.get_string_width(dates) if dates else 0
        y0, page0 = pdf.get_y(), pdf.page_no()
        self.font("B")
        self.text_at(e.get("heading") or "", pdf.l_margin, pdf.epw - dates_w - (4 if dates else 0), self.lh())
        y1 = pdf.get_y()
        if dates and pdf.page_no() == page0:
            self.font("")
            pdf.set_xy(pdf.w - pdf.r_margin - dates_w, y0)
            pdf.cell(dates_w, self.lh(), dates, align="R")
            pdf.set_xy(pdf.l_margin, y1)
        sub = ", ".join(x for x in (e.get("subheading"), e.get("location")) if x)
        if sub:
            self.block(sub, style="I", color=self.SECONDARY)
        if e.get("summary"):
            self.block(e["summary"])
        for b in e.get("bullets", []):
            self.bullet(item_text(b))
        pdf.ln(1.4)

    def group(self, g):
        pdf, h = self.pdf, self.lh()
        self.ensure(h * 1.5)
        if g.get("label"):
            self.font("B")
            pdf.write(h, self.t(g["label"] + ": "))
        self.font("")
        pdf.write(h, self.t(", ".join(g.get("items", []))))
        pdf.ln(h + 0.6)

    def render(self, path: Path) -> int:
        pdf, header, meta = self.pdf, self.cv["header"], self.cv.get("meta", {})
        pdf.set_title(meta.get("title") or f"{header['name']} - CV")
        pdf.set_author(header["name"])
        pdf.set_lang(self.lang)
        pdf.add_page()
        self.header()
        first_block_lines = {"text": 2, "entries": 4, "groups": 1.5, "list": 1.2}
        for sec in self.cv["sections"]:
            kind = sec["kind"]
            if not section_has_content(sec):
                continue
            self.section_title(sec["title"], first_block_lines[kind])
            if kind == "text":
                for para in split_paragraphs(sec.get("text", "")):
                    self.block(para, after=1.0)
            elif kind == "entries":
                for e in sec.get("entries", []):
                    self.entry(e)
            elif kind == "groups":
                for g in sec.get("groups", []):
                    self.group(g)
            elif kind == "list":
                for item in sec.get("items", []):
                    self.bullet(item_text(item))
        pdf.output(str(path))
        missing = {c for font in pdf.fonts.values() for c in getattr(font, "missing_glyphs", [])}
        self.missing_glyphs = [chr(c) for c in sorted(missing)]
        return pdf.page_no()


# ---------------------------------------------------------------------------
# DOCX rendering


def render_docx(cv, path: Path):
    from docx import Document
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Mm, Pt, RGBColor

    meta, header = cv.get("meta", {}), cv["header"]
    lang = meta.get("language") or "en"
    base = float(meta.get("fontSize") or 10.5)
    latin = "Arial"
    east_asian = DOCX_EAST_ASIAN_FONT[preferred_han_script(lang, meta.get("market") or "")]
    primary, secondary = RGBColor(0x1F, 0x23, 0x28), RGBColor(0x57, 0x60, 0x6A)

    doc = Document()
    section = doc.sections[0]
    if str(meta.get("paper", "A4")).lower() == "letter":
        section.page_width, section.page_height, side = Inches(8.5), Inches(11), Inches(0.7)
    else:
        section.page_width, section.page_height, side = Mm(210), Mm(297), Mm(17)
    section.left_margin = section.right_margin = side
    section.top_margin = section.bottom_margin = Mm(15)
    text_width = section.page_width - section.left_margin - section.right_margin

    def fix_fonts(rpr):
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rfonts.attrib.pop(qn(attr), None)
        for attr, value in (("w:ascii", latin), ("w:hAnsi", latin), ("w:cs", latin), ("w:eastAsia", east_asian)):
            rfonts.set(qn(attr), value)

    def style_font(style, size, bold=None, color=primary):
        style.font.size = Pt(size)
        style.font.bold = bold
        rpr = style.element.get_or_add_rPr()
        fix_fonts(rpr)
        old = rpr.find(qn("w:color"))
        if old is not None:
            rpr.remove(old)
        style.font.color.rgb = color

    defaults = doc.styles.element.find(qn("w:docDefaults"))
    default_rpr = defaults.find(qn("w:rPrDefault")).find(qn("w:rPr")) if defaults is not None else None
    if default_rpr is not None:
        fix_fonts(default_rpr)

    normal = doc.styles["Normal"]
    style_font(normal, base)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.08

    heading = doc.styles["Heading 1"]
    style_font(heading, base + 1.5, bold=True)
    heading.paragraph_format.space_before = Pt(10)
    heading.paragraph_format.space_after = Pt(3)
    heading.paragraph_format.keep_with_next = True
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for attr, value in (("w:val", "single"), ("w:sz", "4"), ("w:space", "1"), ("w:color", "A0A0A0")):
        bottom.set(qn(attr), value)
    border.append(bottom)
    heading.element.get_or_add_pPr().insert_element_before(
        border, "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
        "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
        "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
        "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle",
        "w:rPr", "w:sectPr", "w:pPrChange")
    style_font(doc.styles["List Bullet"], base)

    def run(p, text, size=None, bold=False, italic=False, color=None):
        r = p.add_run(text)
        r.bold = bold or None
        r.italic = italic or None
        if size:
            r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color
        return r

    def para(text="", size=None, bold=False, italic=False, color=None, before=0, after=2, keep=False, style=None):
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        if keep:
            p.paragraph_format.keep_with_next = True
        if text:
            run(p, text, size, bold, italic, color)
        return p

    para(header["name"], size=base + 7.5, bold=True, after=1)
    if header.get("headline"):
        para(header["headline"], size=base + 0.5, color=secondary, after=1)
    para("  |  ".join(header.get("contact", [])), size=base - 1, after=1)
    if header.get("details"):
        para("  |  ".join(header["details"]), size=base - 1, color=secondary, after=1)

    for sec in cv["sections"]:
        if not section_has_content(sec):
            continue
        doc.add_paragraph(sec["title"], style="Heading 1")
        kind = sec["kind"]
        if kind == "text":
            for chunk in split_paragraphs(sec.get("text", "")):
                para(chunk, after=3)
        elif kind == "entries":
            for e in sec.get("entries", []):
                p = para(before=4, after=0, keep=True)
                p.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
                run(p, e.get("heading", ""), bold=True)
                if e.get("dates"):
                    run(p, "\t" + e["dates"])
                sub = ", ".join(x for x in (e.get("subheading"), e.get("location")) if x)
                if sub:
                    para(sub, italic=True, color=secondary, after=1, keep=bool(e.get("bullets") or e.get("summary")))
                if e.get("summary"):
                    para(e["summary"], after=1)
                for b in e.get("bullets", []):
                    para(item_text(b), style="List Bullet", after=1)
        elif kind == "groups":
            for g in sec.get("groups", []):
                p = para(after=1)
                if g.get("label"):
                    run(p, g["label"] + ": ", bold=True)
                run(p, ", ".join(g.get("items", [])))
        elif kind == "list":
            for item in sec.get("items", []):
                para(item_text(item), style="List Bullet", after=1)

    # Replace the python-docx template's metadata (author "python-docx", 2013 dates).
    now = dt.datetime.now(dt.timezone.utc).replace(tzinfo=None, microsecond=0)
    props = doc.core_properties
    props.title = meta.get("title") or f"{header['name']} - CV"
    props.author = props.last_modified_by = header["name"]
    props.created = props.modified = now
    props.revision = 1
    props.language = lang
    props.comments = props.keywords = props.subject = props.category = ""
    doc.save(str(path))


# ---------------------------------------------------------------------------
# Job posting fetch


class PageParser(HTMLParser):
    # <head> is not skipped as a whole: its end tag is optional, and <title> is captured separately.
    SKIP = {"script", "style", "noscript", "template", "svg", "iframe"}
    BLOCK = {"p", "div", "br", "ul", "ol", "h1", "h2", "h3", "h4", "h5", "h6", "tr", "section", "article",
             "header", "footer", "table", "dd", "dt", "hr", "blockquote", "pre"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts, self.ld_json, self.comments, self.meta = [], [], [], {}
        self.title = ""
        self._skip = 0
        self._in_ld = self._in_title = False
        self._buf = []

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag == "script" and (a.get("type") or "").lower() == "application/ld+json":
            self._in_ld, self._buf = True, []
            return
        if tag == "meta":
            key = a.get("property") or a.get("name")
            if key in ("og:title", "og:description", "og:site_name", "description") and a.get("content"):
                self.meta[key] = a["content"]
        if tag == "title":
            self._in_title = True
        if tag in self.SKIP:
            self._skip += 1
        elif tag == "li":
            self.parts.append("\n- ")
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag == "script" and self._in_ld:
            self.ld_json.append("".join(self._buf))
            self._in_ld = False
            return
        if tag == "title":
            self._in_title = False
        if tag in self.SKIP and self._skip:
            self._skip -= 1
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if self._in_ld:
            self._buf.append(data)
        elif self._in_title:
            self.title += data
        elif not self._skip:
            self.parts.append(data)

    def handle_comment(self, data):
        self.comments.append(data)


def clean_text(text: str) -> str:
    lines = [re.sub(r"[ \t\u00a0]+", " ", line).strip() for line in text.replace("\r", "").split("\n")]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def html_to_text(fragment: str) -> str:
    if "&lt;" in fragment and "<" not in fragment:
        fragment = html.unescape(fragment)
    parser = PageParser()
    parser.feed(fragment)
    parser.close()
    return clean_text("".join(parser.parts))


def find_job_posting(blocks):
    def walk(node):
        if isinstance(node, list):
            for n in node:
                yield from walk(n)
        elif isinstance(node, dict):
            yield node
            if "@graph" in node:
                yield from walk(node["@graph"])

    for raw in blocks:
        try:
            data = json.loads(raw.strip(), strict=False)
        except json.JSONDecodeError:
            continue
        for node in walk(data):
            types = node.get("@type")
            if "JobPosting" in (types if isinstance(types, list) else [types]):
                return node
    return None


def posting_fields(posting):
    org = posting.get("hiringOrganization")
    company = org.get("name") if isinstance(org, dict) else org
    locations = posting.get("jobLocation") or []
    places = []
    for loc in locations if isinstance(locations, list) else [locations]:
        address = loc.get("address") if isinstance(loc, dict) else None
        if isinstance(address, dict):
            country = address.get("addressCountry")
            country = country.get("name") if isinstance(country, dict) else country
            places.append(", ".join(str(x) for x in (address.get("addressLocality"),
                                                      address.get("addressRegion"), country) if x))
        elif isinstance(address, str):
            places.append(address)
    location = "; ".join(dict.fromkeys(p for p in places if p))
    if posting.get("jobLocationType") == "TELECOMMUTE":
        location = (location + " (remote)").strip()
    return {
        "title": html.unescape(posting.get("title") or ""),
        "company": html.unescape(company or ""),
        "location": location,
        "employmentType": posting.get("employmentType"),
        "datePosted": posting.get("datePosted"),
        "validThrough": posting.get("validThrough"),
        "salary": salary_text(posting.get("baseSalary")),
    }


def salary_text(salary) -> str:
    if not isinstance(salary, dict):
        return str(salary or "")
    value = salary.get("value")
    if isinstance(value, dict):
        low, high = value.get("minValue"), value.get("maxValue")
        amount = f"{low}–{high}" if low is not None and high is not None else str(value.get("value") or low or high or "")
        unit = value.get("unitText") or ""
    else:
        amount, unit = str(value or ""), ""
    return " ".join(x for x in (salary.get("currency") or "", amount, f"per {unit.lower()}" if unit else "") if x)


def sniff_charset(raw: bytes):
    head = raw[:4096].decode("ascii", errors="ignore")
    m = re.search(r"""<meta[^>]+charset\s*=\s*["']?([\w-]+)""", head, re.I)
    return m.group(1) if m else None


def decode_page(raw: bytes, header_charset):
    for charset in (header_charset, sniff_charset(raw), "utf-8"):
        if not charset:
            continue
        try:
            return raw.decode(charset, errors="replace")
        except LookupError:
            continue
    return raw.decode("utf-8", errors="replace")


INJECTION_PATTERNS = [re.compile(p, re.I) for p in (
    r"ignore (?:all |any )?(?:the )?(?:previous|prior|above|earlier) (?:instructions|prompts?)",
    r"disregard (?:all |any )?(?:the )?(?:previous|prior|above) ",
    r"\bif you are (?:an? )?(?:ai|a\.i\.|llm|large language model|language model|chatgpt|gpt|claude|bot|assistant)\b",
    r"\b(?:ai|llm|chatgpt|gpt|claude|language model)s?\b.{0,40}\b(?:must|should)\s+(?:include|mention|insert|add|write|say)\b",
    r"\bsystem prompt\b",
    r"\binclude the (?:word|phrase)\b",
)]


def scan_injection(text: str):
    hits = []
    for pattern in INJECTION_PATTERNS:
        for m in pattern.finditer(text):
            start, end = max(0, m.start() - 80), min(len(text), m.end() + 80)
            hits.append(re.sub(r"\s+", " ", text[start:end]).strip())
    return list(dict.fromkeys(hits))[:10]


FETCH_HINT = ("Try WebFetch, and if that also fails ask the user to paste the job description text. "
              "Sites behind a login (e.g. LinkedIn) or rendered with JavaScript often cannot be fetched.")
USER_AGENT = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) "
              "Chrome/130.0 Safari/537.36")


def cmd_fetch_jd(args):
    url = args.url.strip()
    if not re.match(r"^https?://", url, re.I):
        die("Expected an http(s) URL.")
    request = urllib.request.Request(url, headers={
        "User-Agent": USER_AGENT,
        "Accept": "text/html,application/xhtml+xml",
        "Accept-Language": "en;q=0.9,*;q=0.5",
    })
    try:
        with urllib.request.urlopen(request, timeout=25) as resp:
            raw = resp.read(5_000_000)
            final_url, status = resp.geturl(), resp.status
            content_type = (resp.headers.get_content_type() or "").lower()
            header_charset = resp.headers.get_content_charset()
    except urllib.error.HTTPError as e:
        emit({"ok": False, "url": url, "status": e.code, "error": f"HTTP {e.code}", "hint": FETCH_HINT})
        return
    except (urllib.error.URLError, OSError, ValueError) as e:
        emit({"ok": False, "url": url, "error": str(e), "hint": FETCH_HINT})
        return

    result = {"url": url, "finalUrl": final_url, "status": status, "contentType": content_type}
    warnings, comments, visible = [], [], ""
    if content_type == "application/pdf" or raw[:5] == b"%PDF-":
        from pypdf import PdfReader

        try:
            text = clean_text("\n\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(raw)).pages))
        except Exception as e:  # damaged or encrypted PDF
            emit({"ok": False, "url": url, "error": f"Could not read the PDF: {e}", "hint": FETCH_HINT})
            return
        result.update(source="pdf")
    elif content_type and not (content_type.startswith("text/") or "html" in content_type or "xml" in content_type):
        emit({"ok": False, "url": url, "contentType": content_type,
              "error": f"Unsupported content type {content_type}.", "hint": FETCH_HINT})
        return
    else:
        parser = PageParser()
        parser.feed(decode_page(raw, header_charset))
        parser.close()
        visible, comments = clean_text("".join(parser.parts)), parser.comments
        posting = find_job_posting(parser.ld_json)
        result["pageTitle"] = clean_text(parser.title) or parser.meta.get("og:title", "")
        if posting:
            fields = posting_fields(posting)
            description = html_to_text(posting.get("description") or "")
            head = "\n".join(x for x in (fields["title"], fields["company"], fields["location"],
                                         fields["salary"] and f"Salary: {fields['salary']}") if x)
            text = f"{head}\n\n{description}".strip()
            # The visible page can hold details the structured data leaves out, such as how to apply.
            result.update(source="json-ld", pageText=visible[:30000], **fields)
        else:
            text = visible
            result.update(source="html")
            if re.search(r"\b(sign in|log in|join now|create an account)\b", visible[:3000], re.I) \
                    and len(text) < 1500:
                warnings.append("The page looks like a login wall.")
            if urllib.parse.urlsplit(final_url)[1:3] != urllib.parse.urlsplit(url)[1:3]:
                warnings.append(f"The link redirected to {final_url} and no job posting data was found there; "
                                "the posting may have expired.")

    if len(text) < 400:
        warnings.append("Very little text was found; the page may need JavaScript or a login.")
    if text.count("�") > max(5, len(text) // 200):
        warnings.append("The text contains many undecodable characters; the page encoding was not recognised.")
    result.update(
        ok=not warnings,
        chars=len(text),
        warnings=warnings,
        hint=FETCH_HINT if warnings else None,
        injectionWarnings=scan_injection("\n".join([visible, text] + comments)),
        text=text[:60000],
    )
    emit(result)


# ---------------------------------------------------------------------------
# Commands


def cmd_doctor(args):
    from importlib import metadata

    deps = {}
    for package in ("fpdf2", "python-docx", "pypdf", "jsonschema"):
        try:
            deps[package] = metadata.version(package)
        except metadata.PackageNotFoundError:
            deps[package] = None
    latin = find_latin_font()
    cjk = {}
    for script in NOTO_CJK_INDEX:
        found = find_cjk_font(script)
        cjk[script] = found[0] if found else None
    problems = [f"Missing Python package: {name}" for name, version in deps.items() if version is None]
    notes = []
    if not latin:
        notes.append("No Arial/Liberation/DejaVu font found: PDFs fall back to Helvetica (Latin-1 only). "
                     "Install Liberation Sans or pass --font to render.")
    missing = [script for script, path in cjk.items() if not path]
    if missing:
        notes.append("No system font for " + ", ".join(missing) + ": CVs containing that script need "
                     "--cjk-font or an installed font (e.g. Noto Sans CJK).")
    emit({"ok": not problems, "python": sys.version.split()[0], "platform": platform.platform(),
          "pluginRoot": str(PLUGIN_ROOT), "dependencies": deps,
          "fonts": {"latin": latin["regular"] if latin else None, "cjk": cjk},
          "problems": problems, "notes": notes})


def cmd_init(args):
    data_dir = resolve_data_dir(args.data_dir)
    for sub in ("applications", "sources", "backups"):
        (data_dir / sub).mkdir(parents=True, exist_ok=True)
    profile = data_dir / "profile.json"
    info = {"ok": True, "dataDir": str(data_dir), "profile": str(profile), "profileExists": profile.exists(),
            "applications": len(list((data_dir / "applications").glob("*/application.json")))}
    if profile.exists():
        try:
            data = json.loads(profile.read_text(encoding="utf-8-sig"))
            info.update(profileVersion=data.get("meta", {}).get("version"),
                        profileUpdatedAt=data.get("meta", {}).get("updatedAt"),
                        targetMarkets=[m.get("code") for m in data.get("targetMarkets", [])])
        except json.JSONDecodeError as e:
            info.update(ok=False, error=f"profile.json is not valid JSON: {e}")
    emit(info)


def cmd_extract(args):
    path = Path(args.file).expanduser().resolve()
    text, pages = extract_text(path)
    print(text)
    if path.suffix.lower() == ".pdf" and len(re.sub(r"--- page \d+ ---|\s", "", text)) < 200:
        print("\n[jh] Very little text was extracted: the PDF is probably scanned. "
              "Read the file directly to view the pages.", file=sys.stderr)


def cmd_validate(args):
    data = read_json(args.file)
    errors, warnings = schema_errors(data, args.kind), []
    if not isinstance(data, dict):
        emit({"ok": False, "errors": errors, "warnings": warnings})
        sys.exit(1)
    if args.kind == "profile":
        ids = collect_ids(data)
        errors += [f"Duplicate id: {i}" for i in sorted({i for i in ids if ids.count(i) > 1})]
        if sum(1 for m in data.get("targetMarkets", []) if m.get("primary")) != 1:
            warnings.append("Mark exactly one target market as primary.")
        known = set(ids)
        for skill in data.get("skills", []):
            for ref in skill.get("evidence", []):
                if ref not in known:
                    errors.append(f"Skill {skill.get('name')!r} cites unknown evidence id {ref!r}")
    elif not args.profile:
        warnings.append("Pass --profile to check that every bullet traces back to the profile.")
    elif not errors:
        profile = read_json(args.profile)
        if schema_errors(profile, "profile"):
            die("The profile does not match its schema; run `validate profile` first.")
        trace_errors, trace_warnings = trace_cv(data, profile)
        errors += trace_errors
        warnings += trace_warnings
    emit({"ok": not errors, "errors": errors, "warnings": warnings})
    sys.exit(0 if not errors else 1)


# ---------------------------------------------------------------------------
# Traceability: does every claim in cv.json come from profile.json?

NUMBER_RE = re.compile(r"(?<![A-Za-z\d])\d+(?:[.,]\d+)*")  # digits inside words such as "B2B" are not numbers
YEAR_RE = re.compile(r"(?<!\d)(?:19|20)\d{2}(?!\d)")
NAME_TOKEN_RE = re.compile(r"[A-Za-z0-9][A-Za-z0-9&+#./'-]*[A-Za-z0-9+#]|[A-Za-z]")
CONTENT_WORD_RE = re.compile(r"[a-z]{4,}")
GENERIC_WORDS = {
    "a", "an", "and", "the", "of", "for", "in", "on", "at", "to", "with", "by", "from", "as", "or", "i", "&",
    "present", "current", "today", "now", "jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "sept",
    "oct", "nov", "dec", "january", "february", "march", "april", "june", "july", "august", "september",
    "october", "november", "december",
}
ORG_STOPWORDS = {"ltd", "limited", "inc", "co", "corp", "corporation", "company", "group", "plc", "llc", "gmbh",
                 "the", "of", "and", "&"}
DETAIL_FILLER = {"years", "year", "yrs", "yr", "experience", "exp", "and", "of", "in", "+"}
# Which fields of a cited profile item must show up in the CV entry or list item, and how much of them.
NAME_FIELDS = {
    "experience": [("organization", 1.0), ("title", 0.6)],
    "volunteering": [("organization", 1.0), ("title", 0.6)],
    "education": [("institution", 1.0)],
    "projects": [("name", 0.6)],
    "certifications": [("name", 0.6)],
    "awards": [("name", 0.6)],
    "publications": [("title", 0.6)],
    "languages": [("language", 1.0)],
    "skills": [("name", 1.0)],
    "additional": [],
}
ENTRY_KINDS = {"experience", "volunteering", "education", "projects", "certifications", "awards", "publications",
               "additional"}


def numbers_in(text) -> set:
    out = set()
    for m in NUMBER_RE.finditer(unicodedata.normalize("NFKC", str(text or ""))):
        try:
            out.add(float(m.group().replace(",", "")))
        except ValueError:
            pass
    return out


def fmt_numbers(values) -> str:
    return ", ".join(f"{v:g}" for v in sorted(values))


def ym_date(value):
    m = re.match(r"^(\d{4})(?:-(\d{2}))?$", str(value or ""))
    return dt.date(int(m.group(1)), int(m.group(2) or 1), 1) if m else None


def experience_spans(profile) -> set:
    """Years of experience the CV may state, derived from the position dates."""
    today, spans, starts, ends = dt.date.today(), set(), [], []
    for item in profile.get("experience", []):
        start = ym_date(item.get("start"))
        if not start:
            continue
        end = ym_date(item.get("end")) or today
        starts.append(start)
        ends.append(end)
        years = (end - start).days / 365.25
        spans |= {math.floor(years), math.ceil(years)}
    if starts:
        years = (max(ends) - min(starts)).days / 365.25
        spans |= {math.floor(years), math.ceil(years)}
    return {float(s) for s in spans}


def significant_words(name: str, stopwords) -> list:
    return [w for w in re.findall(r"[^\W_]+", normalize(name)) if w not in stopwords and len(w) > 1]


def name_coverage(name: str, text_norm: str, stopwords=GENERIC_WORDS) -> float:
    """Fraction of a profile name's significant words that appear in the CV text."""
    n = normalize(name)
    if not n or n in text_norm:
        return 1.0
    words = significant_words(name, stopwords)
    if not words:
        return 1.0
    return sum(1 for w in words if keyword_in(text_norm, w)) / len(words)


def term_known(term: str, haystack: str) -> bool:
    t = term.lower().strip(".'")
    if not t or t in GENERIC_WORDS:
        return True
    variants = {t, t + "s", t + "es", t[:-1] if t.endswith("s") else t}
    return any(keyword_in(haystack, v) for v in variants)


def unknown_names(text: str, haystack: str, skip_sentence_start=True) -> list:
    """Capitalized words (companies, schools, tools, degrees) that appear nowhere in the profile."""
    missing = []
    for sentence in re.split(r"(?<=[.!?;:|•])\s+|\n", text or ""):
        for i, token in enumerate(NAME_TOKEN_RE.findall(sentence)):
            if (skip_sentence_start and i == 0) or not any(c.isupper() for c in token):
                continue
            if term_known(token, haystack):
                continue
            parts = [p for p in re.split(r"[/.-]", token) if p]
            if len(parts) > 1 and all(term_known(p, haystack) or not any(c.isupper() for c in p) for p in parts):
                continue
            missing.append(token)
    return list(dict.fromkeys(missing))


def content_overlap(text: str, source: str) -> float:
    """Share of the CV text's content words (or CJK character pairs) found in the cited facts."""
    def tokens(s):
        s = normalize(s)
        cjk = re.sub(f"[^{CJK_CHARS}]", "", s)
        return set(CONTENT_WORD_RE.findall(s)) | {cjk[i:i + 2] for i in range(len(cjk) - 1)}

    words = tokens(text)
    return 1.0 if len(words) < 4 else len(words & tokens(source)) / len(words)


def profile_index(profile):
    """Return ({top-level id: (section, item)}, {any id: its top-level id}, {any id: its object})."""
    top, owner, nodes = {}, {}, {}

    def walk(obj, top_id):
        if isinstance(obj, dict):
            if isinstance(obj.get("id"), str):
                owner[obj["id"]] = top_id
                nodes[obj["id"]] = obj
            for value in obj.values():
                walk(value, top_id)
        elif isinstance(obj, list):
            for value in obj:
                walk(value, top_id)

    for section, value in profile.items():
        if isinstance(value, list):
            for item in value:
                if isinstance(item, dict) and isinstance(item.get("id"), str):
                    top[item["id"]] = (section, item)
                    walk(item, item["id"])
    return top, owner, nodes


def scalar_text(item: dict) -> str:
    """An item's own fields, without nested lists such as other achievements."""
    return " ".join(str(v) for v in item.values() if isinstance(v, (str, int, float)) and not isinstance(v, bool))


def group_candidates(profile, nodes):
    """(name, detail text) pairs that a Skills or Languages group item may refer to."""
    out = []
    for skill in profile.get("skills", []):
        years = f"{skill['years']:g} years" if isinstance(skill.get("years"), (int, float)) else ""
        evidence = " ".join(" ".join(iter_strings(nodes[e])) for e in skill.get("evidence", []) if e in nodes)
        out.append((skill.get("name", ""), " ".join(filter(None, (skill.get("level"), skill.get("category"),
                                                                 years, evidence)))))
    for lang in profile.get("languages", []):
        out.append((lang.get("language", ""), lang.get("proficiency", "")))
    for cert in profile.get("certifications", []):
        out.append((cert.get("name", ""), " ".join(filter(None, (cert.get("issuer"), cert.get("date"))))))
    for node in nodes.values():
        for tag in node.get("skills", []) if isinstance(node.get("skills"), list) else []:
            out.append((tag, " ".join(iter_strings(node))))
    return out


def name_matches(base: str, name: str) -> bool:
    b, n = normalize(base), normalize(name)
    if not b or not n:
        return False
    parts = {normalize(p) for p in re.split(r"\s*[/,&()]\s*", name) if p.strip()}
    return any(b == v or b == v + "s" or b + "s" == v for v in {n} | parts)


def detail_matches(detail: str, candidate_detail: str) -> bool:
    """'Python (5 years, advanced)': every number and word in the brackets must be in the profile item."""
    known_words = set(re.findall(r"[^\W\d_]+", normalize(candidate_detail)))
    if not numbers_in(detail) <= numbers_in(candidate_detail):
        return False
    words = [w for w in re.findall(r"[^\W\d_]+", normalize(detail)) if w not in DETAIL_FILLER]
    return all(w in known_words for w in words)


def trace_cv(cv, profile):
    """Check that the names, dates and numbers in cv.json come from the profile facts they cite."""
    errors, warnings = [], []
    top, owner, nodes = profile_index(profile)
    haystack = normalize(" ".join(list(iter_strings(profile)) + [
        re.sub(r"(?<!^)(?=[A-Z])", " ", key) for key in (profile.get("personal", {}).get("regional") or {})]))
    profile_numbers = numbers_in(" ".join(iter_strings(profile))) | experience_spans(profile)
    markets = profile.get("targetMarkets", [])
    profile_lang = next((m.get("cvLanguage") for m in markets if m.get("primary")), None) or \
        (markets[0].get("cvLanguage") if markets else None)
    translated = bool(profile_lang) and \
        normalize(cv["meta"]["language"]).split("-")[0] != normalize(profile_lang).split("-")[0]

    def strict(message):
        # Wording checks cannot work across languages, so they only warn on a translated CV.
        (warnings if translated else errors).append(message)

    def check_free_text(where, text, skip_sentence_start=True):
        extra = numbers_in(text) - profile_numbers
        if extra:
            errors.append(f"{where}: number(s) {fmt_numbers(extra)} do not appear in the profile.")
        names = unknown_names(text, haystack, skip_sentence_start)
        if names:
            strict(f"{where}: {', '.join(names)} not found in the profile.")

    def check_traced(where, item, allowed_tops, need_name):
        if not isinstance(item, dict) or not item.get("source"):
            errors.append(f"{where}: {item_text(item)[:80]!r} has no source ids.")
            return
        refs = item["source"]
        unknown = [r for r in refs if r not in owner]
        if unknown:
            errors.append(f"{where}: unknown source id(s) {', '.join(unknown)}.")
            return
        foreign = [r for r in refs if allowed_tops is not None and owner[r] not in allowed_tops]
        if foreign:
            errors.append(f"{where}: {', '.join(foreign)} belong to another profile item than this entry.")
        text = item_text(item)
        cited = " ".join(" ".join(iter_strings(nodes[r])) + " " + scalar_text(top[owner[r]][1]) for r in refs)
        extra = numbers_in(text) - numbers_in(cited)
        if extra:
            errors.append(f"{where}: number(s) {fmt_numbers(extra)} in {text[:60]!r} are not in the cited facts.")
        names = unknown_names(text, haystack)
        if names:
            strict(f"{where}: {', '.join(names)} not found in the profile.")
        if need_name:
            for ref in (r for r in refs if r in top):
                section, profile_item = top[ref]
                for field, threshold in NAME_FIELDS.get(section, []):
                    if name_coverage(profile_item.get(field, ""), normalize(text)) < threshold:
                        strict(f"{where}: {text[:60]!r} does not match {ref} ({profile_item.get(field)!r}).")
        if not translated and content_overlap(text, cited) < 0.3:
            warnings.append(f"{where}: {text[:60]!r} shares few words with {', '.join(refs)}; check it says "
                            "the same thing.")

    header = cv["header"]
    personal = profile.get("personal", {})
    names = [personal.get(k, "") for k in ("fullName", "displayName", "localName")]
    if not any(normalize(header["name"]) in normalize(n) or normalize(n) in normalize(header["name"])
               for n in names if n):
        strict(f"Header: name {header['name']!r} differs from the profile.")
    if header.get("headline"):
        check_free_text("Header headline", header["headline"], skip_sentence_start=False)
    personal_text = normalize(" ".join(iter_strings(personal)) + " " + " ".join(iter_strings(profile.get("languages", []))))
    for line in header.get("details", []):
        value = line.split(":", 1)[1] if ":" in line else line  # "Notice period: 1 month" -> check "1 month"
        check_free_text(f"Header details {line!r}", value, skip_sentence_start=False)
    for item in header.get("contact", []):
        if numbers_in(item) - numbers_in(personal_text):
            errors.append(f"Header contact {item!r} does not match the profile's personal details.")
        elif re.sub(r"[\s()-]", "", normalize(item)) not in re.sub(r"[\s()-]", "", personal_text):
            warnings.append(f"Header contact {item!r} is not in the profile's personal details.")

    candidates = group_candidates(profile, nodes)
    for sec in cv["sections"]:
        title, kind = sec["title"], sec["kind"]
        if kind == "text":
            check_free_text(title, sec.get("text", ""))
        elif kind == "entries":
            for e in sec.get("entries", []):
                where = f"{title} > {e.get('heading') or '?'}"
                refs = e.get("source") or []
                if not refs:
                    errors.append(f"{where}: the entry has no source ids (cite the profile item, e.g. exp-…).")
                    continue
                bad = [r for r in refs if r not in top]
                if bad:
                    errors.append(f"{where}: {', '.join(bad)} must be ids of top-level profile items "
                                  "(a position, school, project…), not achievements.")
                    continue
                wrong_kind = [r for r in refs if top[r][0] not in ENTRY_KINDS]
                if wrong_kind:
                    errors.append(f"{where}: {', '.join(wrong_kind)} cannot be shown as an entry.")
                label = normalize(" ".join(filter(None, (e.get("heading"), e.get("subheading"),
                                                         e.get("location"), e.get("summary")))))
                for ref in refs:
                    section, item = top[ref]
                    for field, threshold in NAME_FIELDS.get(section, []):
                        if name_coverage(item.get(field, ""), label, ORG_STOPWORDS) < threshold:
                            message = f"{where}: {field} {item.get(field)!r} of {ref} is missing or changed."
                            (warnings.append(message) if len(refs) > 1 else strict(message))
                    if section == "education" and item.get("qualification") and \
                            name_coverage(item["qualification"], label) < 0.5:
                        warnings.append(f"{where}: qualification differs from {item['qualification']!r}.")
                allowed_years = set()
                for ref in refs:
                    item = top[ref][1]
                    for key in ("start", "end", "date"):
                        if ym_date(item.get(key)):
                            allowed_years.add(ym_date(item[key]).year)
                    if "end" in item and item["end"] is None:
                        allowed_years.add(dt.date.today().year)
                shown = {int(y) for y in YEAR_RE.findall(e.get("dates") or "")}
                if shown - allowed_years:
                    errors.append(f"{where}: dates {e.get('dates')!r} do not match the profile "
                                  f"({', '.join(str(y) for y in sorted(allowed_years)) or 'no dates'}).")
                subtree = " ".join(" ".join(iter_strings(top[r][1])) for r in refs)
                if e.get("summary"):
                    extra = numbers_in(e["summary"]) - numbers_in(subtree)
                    if extra:
                        errors.append(f"{where}: number(s) {fmt_numbers(extra)} in the entry summary are not in "
                                      "the cited facts.")
                    names_missing = unknown_names(e["summary"], haystack)
                    if names_missing:
                        strict(f"{where}: {', '.join(names_missing)} not found in the profile.")
                for b in e.get("bullets", []):
                    check_traced(where, b, set(refs), need_name=False)
        elif kind == "list":
            for item in sec.get("items", []):
                check_traced(title, item, None, need_name=True)
        elif kind == "groups":
            for g in sec.get("groups", []):
                for raw in g.get("items", []):
                    m = re.match(r"^(.*?)\s*\((.*)\)\s*$", raw)
                    base, detail = (m.group(1), m.group(2)) if m else (raw, "")
                    matched = [c for c in candidates if name_matches(base, c[0])]
                    if not matched:
                        strict(f"{title}: {raw!r} is not a skill, language or certification in the profile.")
                    elif detail and not any(detail_matches(detail, c[1]) for c in matched):
                        strict(f"{title}: {raw!r}: the detail in brackets does not match the profile.")
    return errors, warnings


def cmd_render(args):
    cv = read_json(args.cv)
    errors = schema_errors(cv, "cv")
    if errors:
        die("cv.json does not match the schema.", details=errors[:20])
    out_dir = Path(args.out_dir).expanduser().resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    basename = safe_name(args.basename or f"{cv['header']['name']}_CV", 120)
    formats = {f.strip().lower() for f in args.formats.split(",") if f.strip()}
    if not formats <= {"pdf", "docx"}:
        die("--formats accepts pdf, docx or pdf,docx.")
    result = {"ok": True, "files": {}, "warnings": []}
    result["warnings"] += [f"Skipped empty section {sec['title']!r}." for sec in cv["sections"]
                           if not section_has_content(sec)]
    if "pdf" in formats:
        writer = PdfCV(cv, args.font, args.cjk_font)
        pdf_path = out_dir / f"{basename}.pdf"
        result["pages"] = writer.render(pdf_path)
        result["files"]["pdf"] = str(pdf_path)
        result["warnings"] += writer.warnings
        if writer.missing_glyphs:
            result["ok"] = False
            result["missingGlyphs"] = writer.missing_glyphs
            result["warnings"].append("No installed font has these characters, so they are blank in the PDF: "
                                      + " ".join(writer.missing_glyphs)
                                      + ". Replace or remove them in cv.json and render again.")
        target = cv["meta"].get("targetPages")
        if target and result["pages"] > target:
            result["warnings"].append(f"The PDF has {result['pages']} pages; the target is {target}.")
    if "docx" in formats:
        docx_path = out_dir / f"{basename}.docx"
        render_docx(cv, docx_path)
        result["files"]["docx"] = str(docx_path)
    record_file = out_dir / "application.json"
    if record_file.exists():
        record = read_json(record_file)
        record.setdefault("files", {})
        for kind, file in result["files"].items():
            record["files"][kind] = Path(file).name
        if "pages" in result:
            record["pages"] = result["pages"]
        record["renderedAt"] = now_iso()
        write_json(record_file, record)
    emit(result)
    sys.exit(0 if result["ok"] else 1)


def load_keywords(args):
    items = []
    if args.keywords_file:
        data = read_json(args.keywords_file)
        if not isinstance(data, list):
            die(f"{args.keywords_file} must be a JSON list of keywords.")
        items += data
    if args.keywords:
        items += [k.strip() for k in args.keywords.split(",") if k.strip()]
    keywords = []
    for k in items:
        if isinstance(k, str):
            keywords.append({"term": k, "aliases": [], "priority": "must"})
        elif isinstance(k, dict) and isinstance(k.get("term"), str):
            aliases = k.get("aliases") if isinstance(k.get("aliases"), list) else []
            keywords.append({"term": k["term"], "aliases": [str(a) for a in aliases],
                             "priority": k.get("priority", "must")})
        else:
            die(f"Invalid keyword entry {k!r}: expected a string or an object with a 'term'.")
    return keywords


def cmd_check(args):
    path = Path(args.file).expanduser().resolve()
    text, pages = extract_text(path)
    plain = re.sub(r"--- page \d+ ---", " ", text)
    norm = normalize(plain)
    found, missing = [], []
    for kw in load_keywords(args):
        hit = next((t for t in [kw["term"]] + kw["aliases"] if keyword_in(norm, t)), None)
        if hit:
            found.append({"term": kw["term"], "matched": hit, "priority": kw["priority"]})
        else:
            missing.append({"term": kw["term"], "priority": kw["priority"]})
    total = len(found) + len(missing)
    must_total = sum(1 for k in found + missing if k["priority"] == "must")
    must_found = sum(1 for k in found if k["priority"] == "must")
    warnings = []
    if len(norm) < 200:
        warnings.append("Almost no text could be extracted: an ATS would see an empty document.")
    if not re.search(r"[^@\s]+@[^@\s]+\.[a-z]{2,}", norm):
        warnings.append("No email address found in the extracted text.")
    if not re.search(r"\+?\d[\d\s().-]{6,}\d", plain):
        warnings.append("No phone number found in the extracted text.")
    if "\ufffd" in plain:
        warnings.append("Some characters did not extract cleanly (replacement characters found).")
    if args.target_pages and pages and pages > args.target_pages:
        warnings.append(f"{pages} pages; the target is {args.target_pages}.")
    emit({
        "ok": not warnings,
        "file": str(path),
        "pages": pages,
        "characters": len(norm),
        "coverage": {"total": total, "found": len(found), "ratio": round(len(found) / total, 2) if total else None,
                     "mustTotal": must_total, "mustFound": must_found},
        "found": found,
        "missing": missing,
        "warnings": warnings,
        "preview": plain.strip()[:1500],
    })


def profile_version(data_dir: Path):
    try:
        return json.loads((data_dir / "profile.json").read_text(encoding="utf-8-sig")).get("meta", {}).get("version")
    except (OSError, json.JSONDecodeError):
        return None


def cmd_app_new(args):
    data_dir = resolve_data_dir(args.data_dir)
    apps = data_dir / "applications"
    apps.mkdir(parents=True, exist_ok=True)
    created = dt.datetime.now().astimezone()
    slug = f"{created:%Y-%m-%d_%H%M}_{safe_name(args.company, 40).lower()}_{safe_name(args.role, 50).lower()}"
    path, n = apps / slug, 2
    while path.exists():
        path, n = apps / f"{slug}-{n}", n + 1
    path.mkdir()
    stamp = created.isoformat(timespec="seconds")
    record = {
        "id": path.name,
        "createdAt": stamp,
        "company": args.company,
        "role": args.role,
        "location": args.location or "",
        "country": args.country or "",
        "jdUrl": args.url or "",
        "jdSummary": "",
        "status": "tailored",
        "statusHistory": [{"status": "tailored", "at": stamp}],
        "keywordCoverage": None,
        "pages": None,
        "profileVersion": profile_version(data_dir),
        "files": {"jd": "jd.md", "keywords": "keywords.json", "cv": "cv.json", "report": "report.md",
                  "pdf": None, "docx": None},
        "notes": "",
    }
    write_json(path / "application.json", record)
    emit({"ok": True, "id": path.name, "dir": str(path)})


SETTABLE_FIELDS = ("company", "role", "location", "country", "jdUrl", "jdSummary", "status", "keywordCoverage",
                   "notes")


def coverage_value(raw):
    """Accept 0.86, "0.86", "86%" or nothing; return a fraction between 0 and 1, or None."""
    text = str(raw).strip().lower()
    if text in ("", "none", "null", "n/a"):
        return None
    try:
        value = float(text.rstrip("%")) / (100 if text.endswith("%") else 1)
    except ValueError:
        die(f"keywordCoverage must be a number such as 0.86 or 86%, not {raw!r}.")
    if not 0 <= value <= 1:
        die(f"keywordCoverage must be between 0 and 1 (or 0% and 100%), not {raw!r}.")
    return round(value, 2)


def cmd_app_set(args):
    record_file = Path(args.dir).expanduser().resolve() / "application.json"
    record = read_json(record_file)
    for pair in args.pairs:
        key, sep, value = pair.partition("=")
        if not sep or not key:
            die(f"Expected key=value, got {pair!r}")
        if key not in SETTABLE_FIELDS:
            die(f"{key!r} cannot be set. Settable fields: {', '.join(SETTABLE_FIELDS)}")
        if key == "keywordCoverage":
            value = coverage_value(value)
        elif key == "status":
            if value not in STATUSES:
                die(f"Unknown status {value!r}. Use one of: {', '.join(STATUSES)}")
            if value != record.get("status"):
                record.setdefault("statusHistory", []).append({"status": value, "at": now_iso()})
        elif key == "jdUrl" and value and not re.match(r"^https?://", value, re.I):
            die("jdUrl must start with http:// or https://")
        elif key == "country":
            value = value.strip().upper()
            if value and not re.fullmatch(r"[A-Z]{2}", value):
                die("country must be a two-letter ISO code such as HK.")
        record[key] = value
    record["updatedAt"] = now_iso()
    write_json(record_file, record)
    emit({"ok": True, "record": record})


def parse_time(value: str):
    try:
        parsed = dt.datetime.fromisoformat(value)
    except (TypeError, ValueError):
        return dt.datetime.min.replace(tzinfo=dt.timezone.utc)
    return parsed if parsed.tzinfo else parsed.astimezone()  # naive timestamps are local time


def record_file_path(folder: Path, files: dict, kind: str):
    """Path of one of a record's files, ignoring names that would point outside its folder."""
    name = files.get(kind) if isinstance(files, dict) else None
    if not isinstance(name, str) or not name or Path(name).name != name or name in (".", ".."):
        return None
    return folder / name


def load_applications(data_dir: Path):
    records = []
    for record_file in (data_dir / "applications").glob("*/application.json"):
        try:
            record = json.loads(record_file.read_text(encoding="utf-8-sig"))
        except (OSError, json.JSONDecodeError):
            continue
        if not isinstance(record, dict):
            continue
        record["dir"] = str(record_file.parent)
        records.append(record)
    records.sort(key=lambda r: parse_time(r.get("createdAt", "")), reverse=True)
    return records


def filter_applications(records, args):
    def has(field, needle):
        return not needle or needle.lower() in str(field or "").lower()

    try:
        since = dt.date.fromisoformat(args.since) if args.since else None
    except ValueError:
        die(f"--since must be a date in YYYY-MM-DD form, not {args.since!r}.")
    out = [r for r in records
           if has(r.get("company"), args.company) and has(r.get("role"), args.role)
           and (not args.status or r.get("status") == args.status)
           and (not since or parse_time(r.get("createdAt", "")).date() >= since)]
    return out[:args.limit] if args.limit else out


def md_cell(value) -> str:
    return str(value or "").replace("|", "\\|").replace("\n", " ").strip()


def md_link(label: str, path: Path) -> str:
    # file:// URIs are percent-encoded, so the link works with spaces and CJK folder names.
    return f"[{label}]({path.resolve().as_uri()})"


def cmd_history_list(args):
    data_dir = resolve_data_dir(args.data_dir)
    all_records = load_applications(data_dir)
    records = filter_applications(all_records, args)
    if args.format == "json":
        emit({"ok": True, "dataDir": str(data_dir), "total": len(all_records), "records": records})
        return
    if not records:
        print(f"No tailored CVs found in {data_dir}." if not all_records else "No records match the filter.")
        return
    lines = ["| # | Date & time | Role | Company | Location | Status | JD | CV |",
             "|---|---|---|---|---|---|---|---|"]
    for i, r in enumerate(records, 1):
        folder = Path(r["dir"])
        files = r.get("files", {})
        when = parse_time(r.get("createdAt", ""))
        jd_file = record_file_path(folder, files, "jd")
        jd = md_link(jd_file.name, jd_file) if jd_file and jd_file.exists() else "—"
        cv_files = [record_file_path(folder, files, kind) for kind in ("pdf", "docx")]
        cv = " · ".join(md_link(f.name, f) for f in cv_files if f and f.exists()) or "—"
        lines.append(f"| {i} | {when:%Y-%m-%d %H:%M} | {md_cell(r.get('role'))} | {md_cell(r.get('company'))} | "
                     f"{md_cell(r.get('location'))} | {md_cell(r.get('status'))} | {jd} | {cv} |")
    print("\n".join(lines))
    print(f"\nShowing {len(records)} of {len(all_records)}. Data folder: {data_dir.as_posix()}")


HTML_PAGE = """<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>CV History</title>
<style>
:root{--bg:#fff;--fg:#1f2328;--muted:#59636e;--line:#d1d9e0;--panel:#f6f8fa;--accent:#0969da}
@media (prefers-color-scheme:dark){:root{--bg:#0d1117;--fg:#e6edf3;--muted:#9198a1;--line:#3d444d;--panel:#151b23;--accent:#4493f8}}
body{margin:0;padding:24px 16px;background:var(--bg);color:var(--fg);font:14px/1.5 system-ui,-apple-system,"Segoe UI",sans-serif}
main{max-width:1200px;margin:0 auto}
h1{font-size:20px;margin:0 0 4px}
.meta{color:var(--muted);margin:0 0 16px}
input{box-sizing:border-box;width:100%;max-width:360px;padding:6px 10px;margin-bottom:12px;border:1px solid var(--line);border-radius:6px;background:var(--bg);color:var(--fg);font:inherit}
.wrap{overflow-x:auto}
table{border-collapse:collapse;width:100%}
th,td{text-align:left;padding:8px;border-bottom:1px solid var(--line);vertical-align:top}
th{font-weight:600;color:var(--muted);white-space:nowrap}
td.nowrap{white-space:nowrap}
td.jd{min-width:280px}
a{color:var(--accent)}
summary{cursor:pointer}
pre{white-space:pre-wrap;font:13px/1.5 ui-monospace,SFMono-Regular,Consolas,monospace;background:var(--panel);padding:12px;border-radius:6px;max-height:480px;overflow:auto}
</style>
</head>
<body>
<main>
<h1>CV History</h1>
<p class="meta">__COUNT__ tailored CVs · generated __GENERATED__</p>
<input id="q" type="search" placeholder="Filter by role, company or status" aria-label="Filter">
<div class="wrap"><table>
<thead><tr><th>#</th><th>Date &amp; time</th><th>Role</th><th>Company</th><th>Location</th><th>Status</th><th>Coverage</th><th>Job description</th><th>CV</th></tr></thead>
<tbody>
__ROWS__
</tbody>
</table></div>
</main>
<script>
document.getElementById("q").addEventListener("input", function (e) {
  var q = e.target.value.toLowerCase();
  document.querySelectorAll("tbody tr").forEach(function (row) {
    row.hidden = row.textContent.toLowerCase().indexOf(q) === -1;
  });
});
</script>
</body>
</html>
"""


def csv_safe(value) -> str:
    """Stop spreadsheet apps from running cell text as a formula."""
    text = "" if value is None else str(value)
    return "'" + text if text[:1] in ("=", "+", "-", "@", "\t", "\r") else text


def coverage_label(value) -> str:
    return f"{round(value * 100)}%" if isinstance(value, (int, float)) and not isinstance(value, bool) else ""


def cmd_history_export(args):
    data_dir = resolve_data_dir(args.data_dir)
    records = load_applications(data_dir)
    if args.format == "csv":
        out = data_dir / "history.csv"
        columns = ["id", "createdAt", "role", "company", "location", "country", "status", "jdUrl", "jdSummary",
                   "keywordCoverage", "pdf", "docx", "dir"]
        with out.open("w", encoding="utf-8-sig", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
            writer.writeheader()
            for r in records:
                files = r.get("files") if isinstance(r.get("files"), dict) else {}
                row = {**r, "pdf": files.get("pdf") or "", "docx": files.get("docx") or ""}
                writer.writerow({key: csv_safe(row.get(key)) for key in columns})
        emit({"ok": True, "file": str(out), "records": len(records)})
        return

    rows = []
    for i, r in enumerate(records, 1):
        folder = Path(r["dir"])
        files = r.get("files") if isinstance(r.get("files"), dict) else {}
        jd_file = record_file_path(folder, files, "jd")
        jd_text = jd_file.read_text(encoding="utf-8", errors="replace") if jd_file and jd_file.exists() \
            else "(jd.md missing)"
        rel = folder.relative_to(data_dir).as_posix()
        cv_files = [(k, record_file_path(folder, files, k)) for k in ("pdf", "docx")]
        links = " · ".join(
            f'<a href="{html.escape(urllib.parse.quote(f"{rel}/{f.name}"))}">{k.upper()}</a>'
            for k, f in cv_files if f) or "—"
        jd_url = str(r.get("jdUrl") or "")
        source = f'<p><a href="{html.escape(jd_url)}">Original posting</a></p>' \
            if re.match(r"^https?://", jd_url, re.I) else ""
        coverage = r.get("keywordCoverage")
        when = parse_time(r.get("createdAt", ""))
        rows.append(
            f"<tr><td>{i}</td><td class=\"nowrap\">{when:%Y-%m-%d %H:%M}</td>"
            f"<td>{html.escape(str(r.get('role') or ''))}</td><td>{html.escape(str(r.get('company') or ''))}</td>"
            f"<td>{html.escape(str(r.get('location') or ''))}</td><td>{html.escape(str(r.get('status') or ''))}</td>"
            f"<td>{coverage_label(coverage)}</td>"
            f"<td class=\"jd\"><details><summary>{html.escape(str(r.get('jdSummary') or 'Show job description'))}</summary>"
            f"{source}<pre>{html.escape(jd_text)}</pre></details></td><td class=\"nowrap\">{links}</td></tr>")
    page = (HTML_PAGE.replace("__COUNT__", str(len(records)))
            .replace("__GENERATED__", html.escape(dt.datetime.now().strftime("%Y-%m-%d %H:%M")))
            .replace("__ROWS__", "\n".join(rows)))
    out = data_dir / "history.html"
    out.write_text(page, encoding="utf-8")
    emit({"ok": True, "file": str(out), "records": len(records)})


def cmd_profile_backup(args):
    data_dir = resolve_data_dir(args.data_dir)
    source = data_dir / "profile.json"
    if not source.exists():
        die(f"No profile.json in {data_dir}")
    version = profile_version(data_dir) or 0
    dest = data_dir / "backups" / f"profile-v{version}-{dt.datetime.now():%Y%m%d-%H%M%S}.json"
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, dest)
    emit({"ok": True, "backup": str(dest)})


def cmd_source_add(args):
    data_dir = resolve_data_dir(args.data_dir)
    source = Path(args.file).expanduser().resolve()
    if not source.is_file():
        die(f"Not a file: {source}")
    dest = data_dir / "sources" / f"{dt.datetime.now():%Y%m%d-%H%M%S}_{safe_name(source.stem, 80)}{source.suffix.lower()}"
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, dest)
    emit({"ok": True, "copy": str(dest), "relative": dest.relative_to(data_dir).as_posix()})


# ---------------------------------------------------------------------------
# CLI


def build_parser():
    parser = argparse.ArgumentParser(prog="jh", description="Helper CLI for the job-hunting plugin.")
    sub = parser.add_subparsers(dest="command", required=True)

    def add(subparsers, name, func, help_text):
        p = subparsers.add_parser(name, help=help_text)
        p.set_defaults(func=func)
        return p

    add(sub, "doctor", cmd_doctor, "Check Python packages and fonts.")
    p = add(sub, "init", cmd_init, "Create the data folder and report its state.")
    p.add_argument("--data-dir")
    p = add(sub, "extract", cmd_extract, "Print the text of a DOCX, PDF, TXT or MD file.")
    p.add_argument("file")
    p = add(sub, "fetch-jd", cmd_fetch_jd, "Download a job posting and extract its text.")
    p.add_argument("url")
    p = add(sub, "validate", cmd_validate, "Validate profile.json or cv.json.")
    p.add_argument("kind", choices=["profile", "cv"])
    p.add_argument("file")
    p.add_argument("--profile", help="profile.json to trace cv.json bullets against")
    p = add(sub, "render", cmd_render, "Render cv.json to PDF and/or DOCX.")
    p.add_argument("cv")
    p.add_argument("--out-dir", required=True)
    p.add_argument("--basename")
    p.add_argument("--formats", default="pdf")
    p.add_argument("--font", help="TrueType font for Latin text")
    p.add_argument("--cjk-font", help="TrueType/TTC font for CJK text, optionally PATH#INDEX")
    p = add(sub, "check", cmd_check, "ATS round-trip: extract text and report keyword coverage.")
    p.add_argument("file")
    p.add_argument("--keywords", help="comma-separated terms")
    p.add_argument("--keywords-file")
    p.add_argument("--target-pages", type=int)

    app = sub.add_parser("app", help="Application records.").add_subparsers(dest="app_command", required=True)
    p = add(app, "new", cmd_app_new, "Create an application folder.")
    p.add_argument("--data-dir")
    p.add_argument("--company", required=True)
    p.add_argument("--role", required=True)
    p.add_argument("--url")
    p.add_argument("--location")
    p.add_argument("--country")
    p = add(app, "set", cmd_app_set, "Set fields on application.json (key=value): " + ", ".join(SETTABLE_FIELDS))
    p.add_argument("--dir", required=True)
    p.add_argument("pairs", nargs="+")

    history = sub.add_parser("history", help="Tailored CV history.").add_subparsers(dest="history_command",
                                                                                   required=True)
    p = add(history, "list", cmd_history_list, "List tailored CVs, newest first.")
    p.add_argument("--data-dir")
    p.add_argument("--company")
    p.add_argument("--role")
    p.add_argument("--status", choices=STATUSES)
    p.add_argument("--since", help="YYYY-MM-DD")
    p.add_argument("--limit", type=int)
    p.add_argument("--format", choices=["md", "json"], default="md")
    p = add(history, "export", cmd_history_export, "Write history.html or history.csv to the data folder.")
    p.add_argument("--data-dir")
    p.add_argument("--format", choices=["html", "csv"], default="html")

    profile = sub.add_parser("profile", help="Profile maintenance.").add_subparsers(dest="profile_command",
                                                                                   required=True)
    p = add(profile, "backup", cmd_profile_backup, "Copy profile.json into backups/.")
    p.add_argument("--data-dir")

    source = sub.add_parser("source", help="Imported CV files.").add_subparsers(dest="source_command",
                                                                               required=True)
    p = add(source, "add", cmd_source_add, "Copy an imported CV into sources/.")
    p.add_argument("--data-dir")
    p.add_argument("file")
    return parser


def main(argv=None):
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")
        except (AttributeError, ValueError):
            pass
    args = build_parser().parse_args(argv)
    try:
        args.func(args)
    except Exception as e:  # report unexpected failures as JSON, like every other error
        emit({"ok": False, "error": f"{type(e).__name__}: {e}"})
        sys.exit(1)


if __name__ == "__main__":
    main()
